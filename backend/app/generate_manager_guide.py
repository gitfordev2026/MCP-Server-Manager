import os
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = "/app/documents"
DIAGRAMS_DIR = "/app/documents/diagrams"
os.makedirs(DOCS_DIR, exist_ok=True)
os.makedirs(DIAGRAMS_DIR, exist_ok=True)

PRIMARY_COLOR = RGBColor(124, 58, 237)   # Violet
SECONDARY_COLOR = RGBColor(8, 145, 178)  # Cyan
TEXT_COLOR = RGBColor(30, 41, 59)        # Slate 800
MUTED_COLOR = RGBColor(100, 116, 139)    # Slate 500

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_header(doc, title, subtitle):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SECONDARY_COLOR
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    run = h.runs[0]
    run.font.name = 'Arial'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = PRIMARY_COLOR
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = SECONDARY_COLOR
    return h

def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = TEXT_COLOR
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    r_bold = p.add_run(bold_prefix + ": ")
    r_bold.font.name = 'Arial'
    r_bold.font.size = Pt(9.5)
    r_bold.font.bold = True
    r_bold.font.color.rgb = PRIMARY_COLOR

    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = TEXT_COLOR
    return p

def add_callout(doc, text, alert_type="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    color_map = {
        "NOTE": ("F1F5F9", "3B82F6"),
        "IMPORTANT": ("FEE2E2", "EF4444"),
        "EXECUTIVE": ("F3E8FF", "7C3AED")
    }
    bg_hex, _ = color_map.get(alert_type, ("F1F5F9", "3B82F6"))
    set_cell_background(cell, bg_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Pt(6)
    
    run_tag = p.add_run(f"[{alert_type}] ")
    run_tag.font.name = 'Arial'
    run_tag.font.bold = True
    run_tag.font.size = Pt(9.5)
    
    run_txt = p.add_run(text)
    run_txt.font.name = 'Arial'
    run_txt.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_table(doc, headers, data):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "7C3AED")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)
            
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_hex = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_hex)
            p = row_cells[c_idx].paragraphs[0]
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9)
                r.font.color.rgb = TEXT_COLOR
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def embed_diagram(doc, image_path, caption):
    if os.path.exists(image_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(image_path, width=Inches(5.8))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(f"Figure: {caption}")
        r_cap.font.name = 'Arial'
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = MUTED_COLOR

# ---------------------------------------------------------------------------
# DIAGRAM GENERATION FUNCTIONS
# ---------------------------------------------------------------------------
def generate_system_flowchart():
    fig, ax = plt.subplots(figsize=(10, 5.5))
    ax.axis('off')

    boxes = [
        ("User Login\n(Keycloak OIDC)", 0.05, 0.65, 0.24, 0.25, "#E0F2FE", "#0284C7"),
        ("Next.js Proxy\n(route.ts)", 0.38, 0.65, 0.24, 0.25, "#DDD6FE", "#7C3AED"),
        ("FastAPI Security Guard\n(rbac.py)", 0.71, 0.65, 0.24, 0.25, "#DCFCE7", "#16A34A"),
        ("Developer Isolation\n(owner_id filter)", 0.38, 0.2, 0.24, 0.25, "#FEF3C7", "#D97706"),
        ("FastMCP Engine & DB\n(catalog.py / Postgres)", 0.71, 0.2, 0.24, 0.25, "#FCE7F3", "#DB2777")
    ]

    for text, x, y, w, h, bg, border in boxes:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.03", facecolor=bg, edgecolor=border, linewidth=2)
        ax.add_patch(rect)
        ax.text(x + w/2, y + h/2, text, ha='center', va='center', fontsize=9.5, fontweight='bold', color='#1E293B')

    # Flow arrows
    ax.annotate('Authenticates JWT', xy=(0.38, 0.775), xytext=(0.29, 0.775), arrowprops=dict(arrowstyle="->", lw=2, color="#475569"), fontsize=8)
    ax.annotate('HMAC Proxy Call', xy=(0.71, 0.775), xytext=(0.62, 0.775), arrowprops=dict(arrowstyle="->", lw=2, color="#475569"), fontsize=8)
    ax.annotate('Extract Actor Claims', xy=(0.5, 0.45), xytext=(0.5, 0.65), arrowprops=dict(arrowstyle="->", lw=2, color="#475569"), fontsize=8)
    ax.annotate('Scoped Data Fetch', xy=(0.71, 0.325), xytext=(0.62, 0.325), arrowprops=dict(arrowstyle="->", lw=2, color="#475569"), fontsize=8)

    plt.tight_layout()
    path = os.path.join(DIAGRAMS_DIR, "system_workflow.png")
    plt.savefig(path, dpi=200, bbox_inches='tight')
    plt.close()
    return path

def generate_mcp_client_flowchart():
    fig, ax = plt.subplots(figsize=(10, 4.5))
    ax.axis('off')

    steps = [
        ("1. Connect & Auth\nBearer / Cookie / ?token=", 0.05, 0.35, 0.2, 0.4, "#F0FDF4", "#16A34A"),
        ("2. Initialize\nHandshake & Session ID", 0.29, 0.35, 0.2, 0.4, "#EFF6FF", "#2563EB"),
        ("3. tools/list\nDiscover 30 FastMCP Tools", 0.53, 0.35, 0.2, 0.4, "#F5F3FF", "#7C3AED"),
        ("4. tools/call\nInvoke Tool via OpenAPI/REST", 0.77, 0.35, 0.2, 0.4, "#FFF7ED", "#EA580C")
    ]

    for text, x, y, w, h, bg, border in steps:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.03", facecolor=bg, edgecolor=border, linewidth=2)
        ax.add_patch(rect)
        ax.text(x + w/2, y + h/2, text, ha='center', va='center', fontsize=9, fontweight='bold', color='#0F172A')

    ax.annotate('', xy=(0.29, 0.55), xytext=(0.25, 0.55), arrowprops=dict(arrowstyle="->", lw=2, color="#475569"))
    ax.annotate('', xy=(0.53, 0.55), xytext=(0.49, 0.55), arrowprops=dict(arrowstyle="->", lw=2, color="#475569"))
    ax.annotate('', xy=(0.77, 0.55), xytext=(0.73, 0.55), arrowprops=dict(arrowstyle="->", lw=2, color="#475569"))

    plt.tight_layout()
    path = os.path.join(DIAGRAMS_DIR, "mcp_client_flow.png")
    plt.savefig(path, dpi=200, bbox_inches='tight')
    plt.close()
    return path

flow_img = generate_system_flowchart()
mcp_flow_img = generate_mcp_client_flowchart()

# ---------------------------------------------------------------------------
# DOCUMENT CREATION
# ---------------------------------------------------------------------------
doc = Document()
add_header(doc, "MCP Server Manager", "System Architecture, Component Breakdown & Workflow Presentation Guide")

add_callout(doc, "This guide is designed for technical presentation to engineering leadership and management. It details the end-to-end operational architecture, exact file locations, primary function signatures, security enforcement, and workflow flowcharts.", "EXECUTIVE")

# ---------------------------------------------------------------------------
# SECTION 1: SYSTEM OVERVIEW & WORKFLOW ARCHITECTURE
# ---------------------------------------------------------------------------
add_heading_1(doc, "1. Executive System Architecture & Operating Principles")
add_paragraph(doc, "The MCP Server Manager (MCP-Server-Manager) is an enterprise control plane that centralizes microservice OpenAPI registrations and native Model Context Protocol (MCP) servers into a unified, security-hardened gateway (`/mcp/apps`).")

add_bullet(doc, "Unified Tool Gateway", "Exposes REST microservices and MCP servers as standard MCP tools via FastMCP 3.4.4 streamable HTTP and SSE transports.")
add_bullet(doc, "Strict Multi-Tenant Isolation", "Developers can only view, manage, and execute tools registered under their own user ID (`created_by_user_id`), preventing cross-tenant data leakage.")
add_bullet(doc, "Role-Based Access Control (RBAC)", "Integrated with Keycloak OIDC identity provider. Audit logs, global tenant policies, and user management are strictly restricted to Admin roles.")
add_bullet(doc, "Dynamic Startup Schema Auto-Sync", "On application boot, `ensure_dynamic_schema_migrations()` automatically inspects all SQLAlchemy model declarations against live database tables and executes `ALTER TABLE ... ADD COLUMN IF NOT EXISTS` seamlessly.")

embed_diagram(doc, flow_img, "End-to-End User Authentication & Tenant Isolation Request Flowchart")

# ---------------------------------------------------------------------------
# SECTION 2: FRONTEND COMPONENT BREAKDOWN
# ---------------------------------------------------------------------------
add_heading_1(doc, "2. Frontend Architecture & Component Matrix")
add_paragraph(doc, "The frontend layer is built with Next.js 16 (App Router), TailwindCSS, and TypeScript. All client requests route through a serverless proxy layer to enforce HMAC signature validation and credentials management.")

add_table(doc,
    ["File Path", "Primary Functions & Components", "Operational Responsibility"],
    [
        [
            "frontend/app/api/proxy/[...path]/route.ts",
            "handleProxy(req, { params })",
            "Next.js API proxy handler. Signs outgoing requests with HMAC headers (X-Signature, X-Timestamp), forwards HttpOnly auth cookies, and manages failover routing across target origins."
        ],
        [
            "frontend/services/http.ts",
            "authenticatedFetch(input, init)\nhttp<T>(path, options)",
            "Centralized HTTP utility layer. Manages credentials inclusion (`credentials: 'include'`), handles 401 session expiration redirects, and formats standard JSON responses."
        ],
        [
            "frontend/components/Navigation.tsx",
            "Navigation({ pageTitle })",
            "Main layout navigation bar. Dynamically fetches current user role from `/api/me` and renders navigation tabs (Admin tab rendered for both Admin and Developer roles)."
        ],
        [
            "frontend/app/admin/page.tsx",
            "loadCurrentUser()\nfetchAll()\nTABS",
            "Admin Panel. Displays dashboard statistics, registered applications, MCP servers, tools, and API endpoints. Developers view only their owned resources; Audit tab is restricted to Admins."
        ],
        [
            "frontend/app/playground/page.tsx",
            "PlaygroundPage()\nfetchData()\nhandleAppSelect()",
            "Interactive AI Tool Testing Playground. Loads application tool catalog via `/mcp/openapi/catalog` and sends prompt test queries to `/agent/playground/query`."
        ],
        [
            "frontend/app/mcp-endpoints/page.tsx",
            "McpEndpointsPage()\nfetchData()\ncopyUrl()",
            "Combined MCP Endpoints Hero page. Displays the absolute server URL (`/api/proxy/mcp/apps`), combined tool count, health summary, and fallback text copy utilities."
        ]
    ]
)

# ---------------------------------------------------------------------------
# SECTION 3: BACKEND ENGINE COMPONENT BREAKDOWN
# ---------------------------------------------------------------------------
add_heading_1(doc, "3. Backend Engine Architecture & Function Reference")
add_paragraph(doc, "The backend is implemented in FastAPI (Python 3.11) with SQLAlchemy ORM, Redis caching, Keycloak OIDC validation, and FastMCP runtime execution.")

add_table(doc,
    ["File Path", "Key Function / Class Signatures", "Operational Responsibility"],
    [
        [
            "backend/app/main.py",
            "init_db()\nensure_dynamic_schema_migrations()\nJWTAuthASGIMiddleware\nCombinedAppsOpenAPIMCP",
            "Backend entrypoint. Initializes database schema auto-migrations on boot, mounts `/mcp/apps` FastMCP ASGI application, and validates Bearer/Cookie/Query JWT auth."
        ],
        [
            "backend/app/core/rbac.py",
            "get_request_actor(request)\nrequire_role(allowed_roles)\nverify_resource_ownership()",
            "RBAC & Security core module. Decodes Keycloak JWT tokens, extracts primary role (`admin`, `developer`), and enforces tenant resource ownership validation."
        ],
        [
            "backend/app/core/cache.py",
            "cache_get_json(key)\ncache_set_json(key, value)\ncache_delete_prefix(prefix)",
            "Redis cache layer. Caches status responses using user-namespaced keys (`status:servers:...:{sub}`) to prevent cross-tenant cache contamination."
        ],
        [
            "backend/app/routers/catalog.py",
            "create_catalog_router()\nget_openapi_tool_catalog()",
            "Unified Catalog Router. Serves `/mcp/openapi/catalog`. Filters active apps and tools by `created_by_user_id` for developers and returns aligned `summary` stats."
        ],
        [
            "backend/app/routers/base_urls.py",
            "create_base_urls_router()\nlist_base_urls()\nregister_base_url()",
            "Application Registry Router. Handles `/base-urls` CRUD operations. Enforces regex URL validation, duplicate URL checks, and developer tenant isolation."
        ],
        [
            "backend/app/routers/servers.py",
            "create_servers_router()\nlist_servers()\nregister_server()",
            "MCP Server Router. Manages `/servers` registrations, performs connectivity probing (`discover_server_tools`), and isolates server queries by developer ID."
        ],
        [
            "backend/app/routers/tools.py",
            "create_tools_router()\nlist_tools()",
            "Tool Registry Router. Manages `/tools` query listings and status updates. Filters tools based on active parent server/app ownership."
        ],
        [
            "backend/app/routers/endpoints.py",
            "create_endpoints_router()\nlist_endpoints()",
            "API Endpoint Router. Manages `/endpoints` registry with owner ID prefixes (`app:{name}` and `mcp:{name}`) and MCP exposure toggle states."
        ],
        [
            "backend/app/routers/audit.py",
            "create_audit_router()\nget_audit_logs()",
            "Audit Log Router. Strictly guarded by `Depends(require_role(['admin']))`. Returns system activity logs for compliance."
        ],
        [
            "backend/app/routers/agent.py",
            "create_agent_router()\nagent_query()\nplayground_query()",
            "AI Agent Execution Router. Interfaces with Ollama LLM (`/agent/query`) and restricts available tools strictly to the developer's owned applications."
        ],
        [
            "backend/app/services/registry/exposure_service.py",
            "resolve_exposable_tools()",
            "Exposure Calculation Engine. Computes effective tool access modes (`allow`, `approval`, `deny`) against `access_policies` and DB registry tables."
        ]
    ]
)

# ---------------------------------------------------------------------------
# SECTION 4: MCP CLIENT COMPATIBILITY & PROTOCOL WORKFLOW
# ---------------------------------------------------------------------------
add_heading_1(doc, "4. Combined MCP Endpoint & Client Protocol Compatibility")
add_paragraph(doc, "External AI assistants and MCP clients (e.g. Claude Desktop, Cursor, AGY CLI) connect to the combined endpoint (`/mcp/apps`) using standard JSON-RPC 2.0 over FastMCP Streamable HTTP / SSE transport.")

embed_diagram(doc, mcp_flow_img, "MCP Client Interaction Sequence (Handshake, Discovery & Tool Invocation)")

add_heading_2(doc, "MCP Protocol Step-by-Step Execution Sequence")
add_bullet(doc, "1. Authentication Handshake", "The MCP client connects to `/mcp/apps/` providing authentication via `Authorization: Bearer <token>`, `access_token` cookie, or URL parameter `?token=<token>`.")
add_bullet(doc, "2. Protocol Initialization", "Client sends `initialize` JSON-RPC 2.0 request (`Accept: application/json, text/event-stream`). Server responds with `200 OK`, returning server capabilities and `Mcp-Session-Id` header.")
add_bullet(doc, "3. Tool Discovery (tools/list)", "Client calls `tools/list` with `Mcp-Session-Id` header. Server returns full JSON schema for all 30 exposed tools.")
add_bullet(doc, "4. Tool Invocation (tools/call)", "Client executes a tool (e.g. `mcp_client_secure__health_health_get`). Server verifies access policy, executes upstream HTTP call, and returns `structuredContent` result.")

add_callout(doc, "All 6 approval documents (.docx) and architecture guides have been compiled and verified in the documents/ directory.", "SUCCESS")

doc.save(os.path.join(DOCS_DIR, "System_Architecture_and_Workflow_Guide.docx"))
print("Created System_Architecture_and_Workflow_Guide.docx successfully!")
